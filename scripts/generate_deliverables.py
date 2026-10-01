"""Generate a dataset-backed PDF report and PowerPoint in the project root."""

from __future__ import annotations

import argparse
import logging
import sys
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.config import PROJECT_ROOT
from src.task3_rating_classification import ClassificationResult, train_rating_models
from src.task4_segmentation import SegmentationResult, segment_content
from src.task5_forecasting import ForecastResult, forecast_releases
from src.task6_analytics_engine import (
    AnalyticsBundle,
    generate_analytics,
    save_plotly_png,
)
from src.utils import load_default_dataset

LOGGER = logging.getLogger(__name__)
NAVY = "#1E293B"
BLUE = "#2563EB"
GREEN = "#10B981"
AMBER = "#F59E0B"
RED = "#EF4444"
PALETTE = [BLUE, GREEN, AMBER, RED, "#0F766E", "#64748B"]
REPORT_PATH = PROJECT_ROOT / "Netflix_Project_Report.pdf"
PRESENTATION_PATH = PROJECT_ROOT / "Netflix_Project_Presentation.pptx"


def build_charts(
    analytics: AnalyticsBundle,
    classification: ClassificationResult,
    segmentation: SegmentationResult,
    forecasting: ForecastResult,
) -> dict[str, Path]:
    """Render the figures used by both deliverables as PNG files."""
    charts: dict[str, Path] = {}
    genre_figure = px.bar(
        analytics.genres.head(12).sort_values("Titles"),
        x="Titles",
        y="Genre",
        orientation="h",
        title="Most represented genres",
        color="Titles",
        color_continuous_scale=["#BFDBFE", BLUE],
    )
    charts["genres"] = save_plotly_png(genre_figure, "report_top_genres.png")

    rating_figure = px.bar(
        analytics.ratings.head(12),
        x="Rating",
        y="Titles",
        title="Rating category distribution",
        color="Rating",
        color_discrete_sequence=PALETTE,
    )
    charts["ratings"] = save_plotly_png(rating_figure, "report_rating_distribution.png")

    metric_columns = ["Accuracy", "Precision", "Recall", "F1 score"]
    metric_long = classification.metrics.melt(
        id_vars="Model", value_vars=metric_columns, var_name="Metric", value_name="Score"
    )
    model_figure = px.bar(
        metric_long,
        x="Model",
        y="Score",
        color="Metric",
        barmode="group",
        title="Rating-model holdout comparison",
        color_discrete_sequence=PALETTE,
    )
    charts["models"] = save_plotly_png(model_figure, "report_rating_model_comparison.png")

    cluster_figure = px.scatter(
        segmentation.titles,
        x="pca_1",
        y="pca_2",
        color=segmentation.titles["cluster"].astype(str),
        symbol="type",
        hover_name="title",
        title="Content segments projected with PCA",
        labels={"color": "Cluster"},
        color_discrete_sequence=PALETTE,
    )
    charts["clusters"] = save_plotly_png(cluster_figure, "report_content_segments.png")

    history = forecasting.monthly_history.rename("Observed").reset_index()
    history.columns = ["Date", "Observed"]
    best_model = str(forecasting.metrics.iloc[0]["Model"])
    future = forecasting.forecasts[
        (forecasting.forecasts["Model"] == best_model) & (forecasting.forecasts["Horizon"] == 12)
    ]
    forecast_figure = go.Figure()
    forecast_figure.add_trace(
        go.Scatter(x=history["Date"], y=history["Observed"], mode="lines", name="Observed")
    )
    forecast_figure.add_trace(
        go.Scatter(x=future["Date"], y=future["Forecast"], mode="lines", name=f"{best_model} forecast", line={"dash": "dash"})
    )
    forecast_figure.update_layout(
        title=f"Catalog additions: observed history and 12-month {best_model} forecast",
        xaxis_title="Month",
        yaxis_title="Titles added",
    )
    charts["forecast"] = save_plotly_png(forecast_figure, "report_release_forecast.png")
    return charts


def _pdf_table(rows: list[list[object]], widths: list[float] | None = None) -> Table:
    table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NAVY)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def _pdf_rows(frame: pd.DataFrame, limit: int = 20, decimals: int = 3) -> list[list[object]]:
    view = frame.head(limit).copy()
    for column in view.select_dtypes(include="number").columns:
        view[column] = view[column].map(lambda value: f"{value:,.{decimals}f}" if pd.notna(value) else "")
    return [view.columns.tolist(), *view.astype(str).values.tolist()]


def create_project_report(
    analytics: AnalyticsBundle,
    classification: ClassificationResult,
    segmentation: SegmentationResult,
    forecasting: ForecastResult,
    charts: dict[str, Path],
) -> None:
    """Write the comprehensive measured-results report to the project root."""
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportSubtitle", parent=styles["Normal"], fontSize=10, leading=14, textColor=colors.HexColor("#475569")))
    styles.add(ParagraphStyle(name="Insight", parent=styles["BodyText"], alignment=TA_LEFT, leftIndent=12, firstLineIndent=-8, leading=14))
    document = SimpleDocTemplate(
        str(REPORT_PATH),
        pagesize=letter,
        rightMargin=0.6 * inch,
        leftMargin=0.6 * inch,
        topMargin=0.58 * inch,
        bottomMargin=0.58 * inch,
        title="Netflix Catalog Intelligence Project Report",
        author="Netflix Catalog Intelligence",
    )
    story: list[object] = [
        Paragraph("Netflix Catalog Intelligence", styles["Title"]),
        Paragraph("Dataset-backed project report · generated " + datetime.now(timezone.utc).strftime("%B %d, %Y"), styles["ReportSubtitle"]),
        Spacer(1, 14),
        Paragraph("Executive summary", styles["Heading2"]),
    ]
    for insight in analytics.insights:
        story.extend([Paragraph(f"• {insight}", styles["Insight"]), Spacer(1, 4)])
    best_classification = classification.metrics.iloc[0]
    story.extend(
        [
            Spacer(1, 6),
            Paragraph(
                f"The strongest rating classifier on the stratified holdout was {classification.best_model_name} "
                f"(accuracy {best_classification['Accuracy']:.3f}, weighted F1 {best_classification['F1 score']:.3f}). "
                f"Clustering selected K={segmentation.titles['cluster'].nunique()} by silhouette score. "
                f"Forecast evaluation used a chronological holdout of {min(12, max(3, len(forecasting.monthly_history) // 5))} months.",
                styles["BodyText"],
            ),
            PageBreak(),
            Paragraph("1. Dataset overview", styles["Heading1"]),
            Paragraph("All statistics in this report are computed from the supplied Netflix CSV after shared cleaning and duplicate removal.", styles["BodyText"]),
            Spacer(1, 8),
            _pdf_table(_pdf_rows(analytics.summary, limit=30, decimals=0), [3.1 * inch, 3.1 * inch]),
            Spacer(1, 12),
            Paragraph("Missing-value profile", styles["Heading2"]),
            _pdf_table(_pdf_rows(analytics.missing, limit=14, decimals=2), [2.3 * inch, 2.1 * inch, 1.8 * inch]),
            PageBreak(),
            Paragraph("2. Catalog composition", styles["Heading1"]),
            Image(str(charts["genres"]), width=6.6 * inch, height=3.5 * inch),
            Paragraph("Most credited countries", styles["Heading2"]),
            _pdf_table(_pdf_rows(analytics.countries, limit=12, decimals=0), [4.5 * inch, 1.5 * inch]),
            Spacer(1, 10),
            Image(str(charts["ratings"]), width=6.6 * inch, height=3.5 * inch),
            PageBreak(),
            Paragraph("3. Rating classification", styles["Heading1"]),
            Paragraph(
                "The target is the dataset's existing content-rating category. This is not a viewership, satisfaction, or popularity label. "
                f"Evaluation used {classification.train_rows:,} training rows and {classification.test_rows:,} stratified holdout rows.",
                styles["BodyText"],
            ),
            Spacer(1, 8),
            Image(str(charts["models"]), width=6.6 * inch, height=3.5 * inch),
            _pdf_table(_pdf_rows(classification.metrics, limit=10, decimals=3), [1.75 * inch, 1.0 * inch, 1.0 * inch, 1.0 * inch, 1.0 * inch, 0.9 * inch]),
            Spacer(1, 8),
            Paragraph("Top encoded feature-importance signals", styles["Heading2"]),
            _pdf_table(_pdf_rows(classification.feature_importance, limit=12, decimals=4), [4.3 * inch, 1.8 * inch]),
            PageBreak(),
            Paragraph("4. Content segmentation", styles["Heading1"]),
            Paragraph(
                f"KMeans selected {segmentation.titles['cluster'].nunique()} clusters by silhouette score. "
                f"The two-component PCA projection explains {segmentation.pca_variance:.1%} of the encoded feature variance. "
                "Cluster labels are descriptive and should not be treated as official content categories.",
                styles["BodyText"],
            ),
            Spacer(1, 8),
            Image(str(charts["clusters"]), width=6.6 * inch, height=3.5 * inch),
            _pdf_table(_pdf_rows(segmentation.insights, limit=20, decimals=1), [0.65 * inch, 0.55 * inch, 1.0 * inch, 1.55 * inch, 1.55 * inch, 1.15 * inch]),
            Paragraph("Cluster-count diagnostics", styles["Heading2"]),
            _pdf_table(_pdf_rows(segmentation.selection, limit=10, decimals=3), [1.4 * inch, 2.3 * inch, 2.3 * inch]),
            PageBreak(),
            Paragraph("5. Catalog-add forecasting", styles["Heading1"]),
            Paragraph(
                f"The monthly series contains {len(forecasting.monthly_history):,} observed calendar months. "
                "Forecasts model titles added to the catalog, not original production volume or official company plans. "
                "The table uses a chronological holdout and reports RMSE, MAE, and R².",
                styles["BodyText"],
            ),
            Spacer(1, 8),
            Image(str(charts["forecast"]), width=6.6 * inch, height=3.5 * inch),
            _pdf_table(_pdf_rows(forecasting.metrics, limit=10, decimals=3), [2.2 * inch, 1.5 * inch, 1.5 * inch, 1.5 * inch]),
            Spacer(1, 8),
            Paragraph("Forecast insights", styles["Heading2"]),
        ]
    )
    for insight in forecasting.insights:
        story.extend([Paragraph(f"• {insight}", styles["Insight"]), Spacer(1, 4)])
    story.extend(
        [
            Spacer(1, 8),
            Paragraph("Interpretation and limitations", styles["Heading2"]),
            Paragraph(
                "Content ratings are not success outcomes. Forecasts extrapolate a short historical catalog-add series and should be read with their holdout errors; "
                "negative R² means the model did not outperform the holdout mean baseline. Genre growth is descriptive, not causal. "
                "For business decisions, validate against updated licensed data and monitor model drift.",
                styles["BodyText"],
            ),
        ]
    )
    document.build(story)


def _add_text(slide: object, text: str, x: float, y: float, width: float, height: float, size: int = 16, color: str = NAVY, bold: bool = False) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def _add_bullets(slide: object, lines: list[str], x: float, y: float, width: float, height: float, size: int = 17) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.level = 0
        paragraph.space_after = Pt(13)
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor.from_string(NAVY.lstrip("#"))


def _add_title(slide: object, title: str, subtitle: str, page: int) -> None:
    _add_text(slide, title, 0.55, 0.35, 12.2, 0.5, 27, NAVY, True)
    _add_text(slide, subtitle, 0.58, 0.91, 12.0, 0.42, 11, "#64748B")
    _add_text(slide, "NETFLIX CATALOG INTELLIGENCE", 0.58, 7.12, 6.0, 0.2, 8, "#64748B", True)
    _add_text(slide, f"{page:02d}", 12.35, 7.08, 0.45, 0.24, 9, BLUE, True)


def _add_metric_card(slide: object, label: str, value: str, x: float, y: float, color: str = BLUE) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.8), Inches(1.18))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string("#F1F5F9".lstrip("#"))
    shape.line.color.rgb = RGBColor.from_string("#CBD5E1".lstrip("#"))
    shape.text_frame.clear()
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    shape.text_frame.margin_left = Inches(0.16)
    shape.text_frame.margin_right = Inches(0.12)
    first = shape.text_frame.paragraphs[0]
    first.text = value
    first.font.name = "Aptos Display"
    first.font.size = Pt(23)
    first.font.bold = True
    first.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    second = shape.text_frame.add_paragraph()
    second.text = label
    second.font.name = "Aptos"
    second.font.size = Pt(10)
    second.font.color.rgb = RGBColor.from_string(NAVY.lstrip("#"))


def _add_chart(slide: object, path: Path, x: float, y: float, width: float, height: float) -> None:
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(width), height=Inches(height))


def create_presentation(
    frame: pd.DataFrame,
    analytics: AnalyticsBundle,
    classification: ClassificationResult,
    segmentation: SegmentationResult,
    forecasting: ForecastResult,
    charts: dict[str, Path],
) -> None:
    """Build a seven-slide presentation with measured results and charts."""
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]
    actual_years = frame["release_year"].dropna()
    added_dates = frame["date_added_dt"].dropna()
    movie_share = frame["type"].eq("Movie").mean() * 100
    best_classification = classification.metrics.iloc[0]
    best_forecast = forecasting.metrics.iloc[0]
    kmeans_best = segmentation.selection.loc[segmentation.selection["Silhouette"].idxmax()]

    slide = presentation.slides.add_slide(blank_layout)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor.from_string(NAVY.lstrip("#"))
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor.from_string(BLUE.lstrip("#"))
    accent.line.fill.background()
    _add_text(slide, "NETFLIX CATALOG", 0.8, 0.85, 6.5, 0.42, 14, "#93C5FD", True)
    _add_text(slide, "Intelligence", 0.76, 1.55, 9.8, 1.15, 46, "#FFFFFF", True)
    _add_text(slide, "Classification · segmentation · forecasting · automated analytics", 0.82, 2.92, 10.8, 0.7, 20, "#CBD5E1")
    _add_text(slide, f"Dataset-backed project presentation  |  {datetime.now(timezone.utc):%B %Y}", 0.84, 6.55, 9.0, 0.35, 12, "#CBD5E1")
    _add_text(slide, "01", 12.0, 6.78, 0.55, 0.3, 10, "#93C5FD", True)

    slide = presentation.slides.add_slide(blank_layout)
    _add_title(slide, "Dataset overview", "Observed catalog snapshot from the supplied Netflix CSV", 2)
    _add_metric_card(slide, "Unique titles", f"{frame['title'].nunique():,}", 0.62, 1.55)
    _add_metric_card(slide, "Movie share", f"{movie_share:.1f}%", 3.67, 1.55, GREEN)
    _add_metric_card(slide, "Ratings represented", str(frame["rating"].nunique()), 6.72, 1.55, AMBER)
    _add_metric_card(slide, "Catalog-add dates", f"{len(added_dates):,}", 9.77, 1.55, RED)
    date_range = f"{added_dates.min():%b %Y} to {added_dates.max():%b %Y}" if not added_dates.empty else "Unavailable"
    release_span = f"{int(actual_years.min())}–{int(actual_years.max())}" if not actual_years.empty else "Unavailable"
    _add_bullets(
        slide,
        [
            f"Original release years: {release_span}",
            f"Observed date_added coverage: {date_range}",
            f"Movies: {frame['type'].eq('Movie').sum():,}  |  TV Shows: {frame['type'].eq('Tv Show').sum():,}",
            "Rating labels describe content suitability, not audience satisfaction or viewing success.",
        ],
        0.82,
        3.25,
        11.4,
        2.6,
        16,
    )

    slide = presentation.slides.add_slide(blank_layout)
    _add_title(slide, "Catalog composition", "Genre and rating patterns are counted from the uploaded rows", 3)
    _add_chart(slide, charts["genres"], 0.55, 1.42, 7.1, 4.95)
    _add_chart(slide, charts["ratings"], 7.82, 1.42, 4.95, 4.95)
    top_genre = analytics.genres.iloc[0] if not analytics.genres.empty else None
    note = f"Largest genre group: {top_genre['Genre']} ({int(top_genre['Titles']):,} tagged titles)." if top_genre is not None else "Genre tags were unavailable."
    _add_text(slide, note, 0.82, 6.47, 11.5, 0.35, 13, NAVY, True)

    slide = presentation.slides.add_slide(blank_layout)
    _add_title(slide, "Rating classification", "Stratified holdout evaluation · best measured model highlighted", 4)
    _add_chart(slide, charts["models"], 0.58, 1.42, 8.0, 4.95)
    _add_metric_card(slide, "Best model", classification.best_model_name, 8.85, 1.62, BLUE)
    _add_metric_card(slide, "Accuracy", f"{best_classification['Accuracy']:.3f}", 8.85, 3.0, GREEN)
    _add_metric_card(slide, "Weighted F1", f"{best_classification['F1 score']:.3f}", 8.85, 4.38, AMBER)
    _add_text(slide, f"{classification.train_rows:,} train rows · {classification.test_rows:,} holdout rows", 8.9, 5.82, 3.9, 0.58, 12, NAVY)
    _add_text(slide, "The target is the existing rating category, not content popularity.", 0.78, 6.48, 11.5, 0.32, 12, "#64748B")

    slide = presentation.slides.add_slide(blank_layout)
    _add_title(slide, "Content segmentation", "KMeans selected by silhouette score · cluster profiles from encoded metadata", 5)
    _add_chart(slide, charts["clusters"], 0.58, 1.42, 7.45, 4.95)
    _add_metric_card(slide, "Selected K", str(int(kmeans_best["Clusters"])), 8.35, 1.72, GREEN)
    _add_metric_card(slide, "Silhouette", f"{kmeans_best['Silhouette']:.3f}", 8.35, 3.12, BLUE)
    _add_metric_card(slide, "PCA variance", f"{segmentation.pca_variance:.1%}", 8.35, 4.52, AMBER)
    _add_text(slide, "Clusters are exploratory groups, not official Netflix categories.", 0.8, 6.5, 11.5, 0.3, 12, "#64748B")

    slide = presentation.slides.add_slide(blank_layout)
    _add_title(slide, "Catalog-add forecasting", "Forecasts extend observed date_added history; they are not company guidance", 6)
    _add_chart(slide, charts["forecast"], 0.58, 1.4, 8.1, 4.95)
    _add_metric_card(slide, "Best holdout model", str(best_forecast["Model"]), 8.92, 1.65, BLUE)
    _add_metric_card(slide, "RMSE", f"{best_forecast['RMSE']:.2f}", 8.92, 3.05, GREEN)
    _add_metric_card(slide, "MAE", f"{best_forecast['MAE']:.2f}", 8.92, 4.45, AMBER)
    _add_text(slide, f"R² {best_forecast['R2']:.3f} · chronological holdout", 8.98, 5.9, 3.7, 0.4, 12, NAVY)

    slide = presentation.slides.add_slide(blank_layout)
    _add_title(slide, "Insights and next steps", "Automated findings with interpretation guardrails", 7)
    _add_bullets(slide, analytics.insights[:6], 0.78, 1.45, 11.8, 3.2, 16)
    _add_text(slide, "Interpretation", 0.82, 4.78, 3.0, 0.42, 19, BLUE, True)
    _add_bullets(
        slide,
        [
            "Negative forecast R² means the model did not beat the holdout mean baseline.",
            "Catalog-add forecasts are extrapolations, not Netflix plans.",
            "Add licensed viewing outcomes before using the analysis to claim content success.",
        ],
        0.82,
        5.18,
        11.5,
        1.3,
        13,
    )

    presentation.save(PRESENTATION_PATH)


def generate() -> tuple[Path, Path]:
    """Load the project dataset, compute every report section, and save outputs."""
    frame = load_default_dataset()
    if frame is None:
        raise FileNotFoundError("No Netflix CSV found. Place one CSV in the project root or data/ and retry.")
    if frame["date_added_dt"].notna().sum() == 0:
        raise ValueError("The presentation/report forecast section requires real date_added values.")
    analytics = generate_analytics(frame)
    classification = train_rating_models(frame, tune=False)
    segmentation = segment_content(frame, max_rows=1500)
    forecasting = forecast_releases(frame)
    charts = build_charts(analytics, classification, segmentation, forecasting)
    create_project_report(analytics, classification, segmentation, forecasting, charts)
    create_presentation(frame, analytics, classification, segmentation, forecasting, charts)
    return REPORT_PATH, PRESENTATION_PATH


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
    report_path, presentation_path = generate()
    print(f"PDF report: {report_path}")
    print(f"PowerPoint: {presentation_path}")


if __name__ == "__main__":
    main()